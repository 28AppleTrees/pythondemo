import os
import pdfplumber
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import logging

# 配置日志记录（可选）
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class PDFToWordConverter:
    """
    一个简单的PDF到Word转换器工具类。
    使用 pdfplumber 提取文本，并用 python-docx 创建 Word 文档。
    """

    def __init__(self, font_name='Arial', font_size=12):
        """
        初始化转换器。

        :param font_name: 输出Word文档的默认字体名称 (e.g., 'Calibri', 'Times New Roman')
        :param font_size: 输出Word文档的默认字体大小 (pt)
        """
        self.font_name = font_name
        self.font_size = font_size

    def convert(self, pdf_path, word_path=None, extract_images=False):
        """
        将指定的PDF文件转换为Word文档 (.docx)。

        :param pdf_path: 源PDF文件的路径。
        :param word_path: 目标Word文件的路径。如果未提供，将基于PDF文件名自动生成。
                          (例如 'example.pdf' -> 'example.docx')
        :param extract_images: 是否尝试提取PDF中的图片 (注意: 当前基础版本不支持，
                               此参数作为预留，不影响当前逻辑)。
        :return: True 如果转换成功，否则 False。
        """
        # --- 输入验证 ---
        if not os.path.exists(pdf_path):
            logger.error(f"错误：源PDF文件 '{pdf_path}' 不存在。")
            return False

        if not pdf_path.lower().endswith('.pdf'):
            logger.warning(f"警告：输入文件 '{pdf_path}' 可能不是PDF格式。")

        # --- 设置输出路径 ---
        if not word_path:
            base_name = os.path.splitext(os.path.basename(pdf_path))[0]
            word_path = f"{base_name}.docx"

        # 确保输出目录存在
        output_dir = os.path.dirname(os.path.abspath(word_path))
        if output_dir and not os.path.exists(output_dir):
            try:
                os.makedirs(output_dir)
                logger.info(f"已创建输出目录: {output_dir}")
            except OSError as e:
                logger.error(f"无法创建输出目录 '{output_dir}': {e}")
                return False

        # --- 核心转换过程 ---
        try:
            # 创建一个新的Word文档对象
            doc = Document()

            # 打开PDF文件
            with pdfplumber.open(pdf_path) as pdf:
                logger.info(f"开始处理PDF文件: {pdf_path} (共 {len(pdf.pages)} 页)")

                for i, page in enumerate(pdf.pages):
                    logger.debug(f"正在处理第 {i + 1} 页...")

                    # 提取页面文本
                    text = page.extract_text()

                    if text:
                        # 在Word文档中添加一个段落
                        paragraph = doc.add_paragraph()

                        # 添加运行(run)，并设置文本内容
                        run = paragraph.add_run(text)

                        # 应用字体样式 (可选，但推荐)
                        run.font.name = self.font_name
                        run.font.size = Pt(self.font_size)
                        # run.font.color.rgb = RGBColor(0, 0, 0) # 黑色

                        # 可以根据需要对齐段落，默认是左对齐
                        # paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

                    else:
                        # 如果页面没有可提取的文本，可以添加一个空行或提示
                        doc.add_paragraph("(此页无文本内容)")
                        logger.debug(f"第 {i + 1} 页未检测到文本。")

                    # --- 图片提取占位符 ---
                    # if extract_images:
                    #     images = page.images
                    #     for img in images:
                    #         # 提取图像数据等操作...
                    #         pass # 当前不实现

            # 保存生成的Word文档
            doc.save(word_path)
            logger.info(f"PDF转换完成。Word文档已保存至: {word_path}")
            return True

        except Exception as e:
            logger.error(f"在转换过程中发生错误: {e}", exc_info=True)
            return False


# --- 示例用法 ---
if __name__ == "__main__":
    # 创建转换器实例，可以自定义字体和字号
    converter = PDFToWordConverter(font_name='Calibri', font_size=11)

    # 指定要转换的PDF文件路径
    input_pdf = "D:\Google\Chrome\download\仪表设计信息模块优化需求.pdf"  # 请替换为你的实际PDF文件路径

    # 可以指定输出的Word文件路径，或者留空让程序自动命名
    output_word = "D:\Google\Chrome\download\converted_output.docx"  # 或者 output_word = None

    # 执行转换
    success = converter.convert(input_pdf, output_word)

    if success:
        print("转换成功！")
    else:
        print("转换失败，请检查日志信息。")